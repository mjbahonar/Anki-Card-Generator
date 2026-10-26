"""Mixed Excel -> Anki integration and exercise privacy/identity checks."""
import copy
import csv
import io
import json
from pathlib import Path
import sqlite3
import tempfile
import tomllib
import unittest
from contextlib import redirect_stdout, closing
from unittest.mock import patch
import zipfile

import pandas as pd
from app.app_config import Config, ConfigError
from app.exercises import read_records, render_exercise, exercise_audio, validate_cards, complete_gap
from app.anki_exporter import generate_anki_package, FIELDS, write_import_csv
from main import run
from app.sources import SourceContext

ROOT = Path(__file__).resolve().parents[1]


class ExerciseTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=ROOT / 'Output')
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        data = tomllib.loads((ROOT / 'config.toml').read_text(encoding='utf-8'))
        data['practice'].update(question_language='en', answer_language='no')
        data['input']['file'] = str(self.directory / 'input.xlsx')
        data['input'].update(header=False, sheet=0, word_column=0, order='excel', max_words=0, deduplicate=True)
        data['output']['directory'] = str(self.directory)
        data['audio']['directory'] = str(self.directory / 'Audio')
        data['style']['file'] = str(ROOT / 'Styles/cards.css')
        data['runtime'].update(keep_awake=False, word_delay_seconds=0)
        data['sources'] = {}
        self.config = Config(ROOT / 'config.toml', data)

    def sheet(self, rows):
        pd.DataFrame(rows).to_excel(self.config.section('input')['file'], header=False, index=False)

    def rendered(self, kind, question='house', answer='hus'):
        return render_exercise(dict(Kind=kind, Words=question, Answer=answer, Sound='[sound:voice.mp3]', Tags=[kind]), self.config)

    def test_optional_columns_and_tags(self):
        self.sheet([['hus'], ['bok']])
        records, skipped = read_records(self.config)
        self.assertFalse(skipped)
        self.assertEqual([r['Kind'] for r in records], ['normal', 'normal'])
        self.assertEqual(records[0]['Tags'], ['normal'])
        self.sheet([['house', 'TYPE', 'A1, daily life,type', 'hus']])
        records, _ = read_records(self.config)
        self.assertEqual(records[0]['Tags'], ['type', 'A1', 'daily_life'])

    def test_shared_media_destination_copies_audio_and_image_keeps_originals(self):
        target = self.directory / 'collection.media'
        self.config.data['media'] = {'anki_directory': str(target)}
        context = SourceContext(self.config)
        self.addCleanup(context.close)
        for name, data in [('voice.mp3', b'ID3' + b'x' * 200), ('photo.png', b'image-data')]:
            original = self.directory / name
            original.write_bytes(data)
            context.register_media(original)
            self.assertEqual(original.read_bytes(), data)
            self.assertEqual((target / name).read_bytes(), data)
            self.assertIn(original.resolve(), context.media)
        self.assertEqual(context.register_audio(self.directory / 'voice.mp3'), '[sound:voice.mp3]')

    def test_gap_alignment_is_explicit_for_ltr_and_rtl(self):
        row = self.rendered('gap', 'Jeg [...] her.', 'bor')
        self.assertIn('style="text-align: left;"', row['FrontField'])
        self.assertIn('style="text-align: left;"', row['BackField'])
        self.config.data['practice']['answer_language'] = 'fa'
        row = self.rendered('gap', 'من [...] هستم.', 'اینجا')
        self.assertIn('style="text-align: right;"', row['FrontField'])
        self.assertIn('style="text-align: right;"', row['BackField'])

    def test_random_order_keeps_answers_and_tags_with_their_questions(self):
        self.sheet([['house', 'type', 'home', 'hus'], ['book', 'type', 'reading', 'bok']])
        self.config.section('input')['order'] = 'random'
        with patch('random.shuffle', side_effect=lambda items: items.reverse()) as shuffle:
            records, _ = read_records(self.config)
        shuffle.assert_called_once()
        self.assertEqual([(r['Words'], r['Answer']) for r in records], [('book', 'bok'), ('house', 'hus')])
        self.assertEqual(records[0]['Tags'], ['type', 'reading'])

    def test_skip_bad_rows_and_allow_optional_listen_meaning(self):
        self.sheet([['house', 'type'], ['I live here.', 'sentence'], ['[…] jente', 'gap'],
                    ['oops', 'typo'], ['hus', 'listen'], ['[…] jente', 'gap', '', 'ei']])
        records, skipped = read_records(self.config)
        self.assertEqual([r['Kind'] for r in records], ['listen', 'gap'])
        self.assertEqual([r['row'] for r in skipped], [1, 2, 3, 4])

    def test_deduplication_keeps_different_types_and_stable_answer_updates(self):
        self.sheet([['house', 'type', '', 'hus'], ['house', 'sentence', '', 'hus'], ['house', 'type', '', 'hus']])
        records, _ = read_records(self.config)
        self.assertEqual(len(records), 2)
        first = self.rendered('type')
        second = self.rendered('type', answer='updated')
        self.assertEqual(first['Identity'], second['Identity'])
        self.assertNotEqual(first['Identity'], self.rendered('sentence')['Identity'])

    def test_instruction_text_is_fixed_and_only_language_name_changes(self):
        for code, name in [('no', 'Norwegian'), ('en', 'English'), ('fa', 'Persian')]:
            self.config.data['practice']['answer_language'] = code
            self.assertIn(f'Say this sentence in {name}.', self.rendered('sentence')['FrontField'])

    def test_persian_direction_language_and_front_does_not_leak_answer(self):
        self.config.data['practice']['answer_language'] = 'fa'
        row = self.rendered('type', answer='خانه')
        self.assertIn('English → Persian', row['FrontField'])
        self.assertIn('Write the equivalent in Persian.', row['FrontField'])
        self.assertNotIn('خانه', row['FrontField'])
        self.assertNotIn('[sound:', row['FrontField'])
        self.assertIn('dir="rtl" lang="fa"', row['BackField'])
        self.assertEqual(row['AnswerDirection'], 'rtl')

    def test_listen_hides_text_and_gap_completes_sentence(self):
        row = self.rendered('listen', question='Jeg bor her.', answer='I live here.')
        self.assertNotIn('Jeg bor her.', row['FrontField'])
        self.assertIn('[sound:voice.mp3]', row['FrontField'])
        self.assertIn('Jeg bor her.', row['BackField'])
        gap = self.rendered('gap', 'Ken [...] 23 år gammel.', 'er')
        from bs4 import BeautifulSoup
        self.assertIn('Ken er 23 år gammel.', BeautifulSoup(gap['BackField'], 'html.parser').get_text())

    def test_multi_gap_replaces_in_order_and_preserves_lines(self):
        question = 'Jeg [...] Ken.\nJeg […] fra England.'
        answer = 'heter\nkommer'
        self.assertEqual(complete_gap(question, answer), 'Jeg heter Ken.\nJeg kommer fra England.')
        row = self.rendered('gap', question, answer)
        self.assertEqual(row['FrontField'].count('class="exercise-gap exercise-gap-entry"'), 2)
        self.assertEqual(row['FrontField'].count('<input'), 2)
        self.assertIn('aria-label="Answer 2"', row['FrontField'])
        self.assertEqual(row['BackField'].count('class="exercise-gap-solution"'), 2)
        self.assertIn('<br>', row['FrontField'])
        self.assertIn('<br>', row['BackField'])
        self.assertIn('exercise-answer exercise-gap-text', row['BackField'])
        self.assertNotIn('heter', row['FrontField'])
        self.assertNotIn('kommer', row['FrontField'])

    def test_speaking_answer_has_no_sample_disclaimer(self):
        row = self.rendered('sentence', 'I live here.', 'Jeg bor her.')
        self.assertIn('Jeg bor her.', row['BackField'])
        self.assertNotIn('Sample answer', row['BackField'])

    def test_multi_gap_invalid_counts_and_empty_answers_are_skipped(self):
        self.sheet([['[...] and […]', 'gap', '', 'one'],
                    ['[...] and [...] and [...]', 'gap', '', 'one\n\nthree'],
                    ['[...] and […]', 'gap', '', 'one\ntwo']])
        records, skipped = read_records(self.config)
        self.assertEqual(len(records), 1)
        self.assertEqual(len(skipped), 2)
        self.assertIn('2 blanks but 1 answer lines', skipped[0]['message'])
        self.assertIn('must not be blank', skipped[1]['message'])

    def test_multi_gap_answers_are_escaped_and_not_regex_replacements(self):
        self.assertEqual(complete_gap('[...] […]', r'\1' + '\n<b>'), r'\1 <b>')
        markup = complete_gap('[...] […]', r'\1' + '\n<b>', markup=True)
        self.assertIn('&lt;b&gt;', markup)
        self.assertNotIn('<b>', markup)

    def test_tts_uses_answer_language_and_reuses_cache(self):
        context = SourceContext(self.config)
        self.addCleanup(context.close)
        with patch('gtts.gTTS') as tts:
            tts.return_value.write_to_fp.side_effect = lambda b: b.write(b'ID3' + b'x' * 200)
            first = exercise_audio('خانه', 'fa', context)
            self.assertEqual(first, exercise_audio('خانه', 'fa', context))
            tts.assert_called_once()
            self.assertEqual(tts.call_args.kwargs['lang'], 'fa')
            exercise_audio('hus', 'nn', context)
            self.assertEqual(tts.call_args.kwargs['lang'], 'no')

    def test_mixed_package_schema_templates_tags_and_media(self):
        self.config.data['practice']['answer_language'] = 'fa'
        rows = [dict(Words='bok', FrontField='bok', Tags=['normal']),
                self.rendered('type', answer='خانه'), self.rendered('sentence'),
                self.rendered('gap', '[...] jente', 'ei'), self.rendered('listen')]
        voice = self.directory / 'voice.mp3'
        voice.write_bytes(b'ID3' + b'x' * 200)
        table_voice = self.directory / 'table.mp3'
        table_voice.write_bytes(b'ID3' + b'y' * 200)
        rows[0]['lexin'] = '<audio controls src="table.mp3"></audio>'
        package = self.directory / 'mixed.apkg'
        generate_anki_package(rows, package, self.config, {voice, table_voice})
        with zipfile.ZipFile(package) as archive:
            database = self.directory / 'collection.anki2'
            database.write_bytes(archive.read('collection.anki2'))
            self.assertIn('voice.mp3', json.loads(archive.read('media')).values())
            self.assertIn('table.mp3', json.loads(archive.read('media')).values())
        with closing(sqlite3.connect(database)) as db:
            models = json.loads(db.execute('select models from col').fetchone()[0])
            notes = db.execute('select mid,tags from notes').fetchall()
            self.assertEqual(len(notes), 5)
            self.assertEqual(len(models), 5)
            self.assertEqual([f['name'] for f in models[str(self.config.section('anki')['model_id'])]['flds']], FIELDS)
            typed = next(m for m in models.values() if m['name'].endswith('type'))
            self.assertIn('{{type:Answer}}', typed['tmpls'][0]['qfmt'])
            self.assertIn('data-answer-direction', typed['tmpls'][0]['qfmt'])
            self.assertIn('data-answer-direction="rtl"', rows[1]['FrontField'])
            self.assertTrue(typed['tmpls'][0]['afmt'].endswith('{{#Info}}{{Info}}{{/Info}}'))
            self.assertTrue(all(tags.strip() for _, tags in notes))

    def test_import_csv_has_note_types_and_space_separated_tags(self):
        row = self.rendered('sentence', 'Hello, friend.', 'Hei, venn.')
        row['Tags'] = ['sentence', 'A1']
        path = self.directory / 'import.csv'
        write_import_csv([row], path, self.config)
        text = path.read_text(encoding='utf-8-sig')
        self.assertIn('#notetype column:1', text)
        parsed = list(csv.reader(line for line in text.splitlines() if not line.startswith('#')))
        self.assertEqual(parsed[0][0], 'Vocabulary Exercise · sentence')
        self.assertEqual(parsed[0][1], 'sentence A1')
        self.assertEqual(parsed[0][4], 'Hello, friend.')
        from app.anki_exporter import note_guid
        self.assertEqual(parsed[0][2], note_guid(row, self.config))
        self.assertIn('#guid column:3', text)

    def test_run_skips_missing_answers_and_no_audio_listen(self):
        self.sheet([['house', 'type', '', 'hus'], ['missing', 'sentence'], ['Jeg bor her.', 'listen']])
        self.config.section('audio')['enabled'] = False
        self.config.section('output')['formats'] = ['json', 'csv', 'apkg', 'html']
        with patch('main.dictionary_content') as dictionary, redirect_stdout(io.StringIO()):
            self.assertEqual(run(self.config), 0)
        dictionary.assert_not_called()
        data = json.loads(next(p for p in self.directory.glob('input_words*.json') if not p.name.endswith('_errors.json')).read_text(encoding='utf-8'))
        self.assertEqual(len(data['words']), 1)
        self.assertEqual(len(data['errors']), 2)
        self.assertTrue(list(self.directory.glob('*_anki.csv')))

    def test_invalid_config_fails_before_network_work(self):
        for update in ({'audio': 'yes'}, {'instruction': '{unknown}'}, {'tag': 'two tags'}):
            data = copy.deepcopy(self.config.data)
            data['cards']['type'].update(update)
            with self.assertRaises(ConfigError):
                validate_cards(data)

    def test_norwegian_audio_cache_follows_explicit_variant_and_tts_support(self):
        self.config.data['language']['source'] = 'nn'
        self.config.data['sources']['lexin'] = {'enabled': True, 'variant': 'bokmal'}
        context = SourceContext(self.config)
        self.addCleanup(context.close)
        self.assertEqual(context.audio_language('lexin'), 'no')
        self.assertEqual(context.audio_language('google_tts'), 'no')
        self.config.data['sources']['lexin']['variant'] = 'both'
        self.assertEqual(context.audio_language('lexin'), 'nn')

    def test_all_invalid_rows_fail_check(self):
        from main import main
        self.sheet([['bad', 'type', '', '']])
        with patch('main.load_config', return_value=self.config), redirect_stdout(io.StringIO()):
            self.assertEqual(main(['--check']), 1)

    def test_typing_multiline_answer_is_skipped(self):
        self.sheet([['house', 'type', '', 'hus\nhuset']])
        records, skipped = read_records(self.config)
        self.assertFalse(records)
        self.assertIn('single-line answer', skipped[0]['message'])

    def test_whole_cell_format_reaches_exercise_front_and_keeps_workbook_unchanged(self):
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill
        self.sheet([['house', 'type', '', 'hus'], ['Say hello', 'sentence', '', 'Hei'], ['bok', 'normal']])
        path = Path(self.config.section('input')['file'])
        workbook = load_workbook(path)
        for coordinate in ('A1', 'A2', 'A3'):
            workbook.active[coordinate].font = Font(bold=True, italic=True, color='FFCC2233')
            workbook.active[coordinate].fill = PatternFill('solid', fgColor='FFFFEE99')
        workbook.save(path)
        workbook.close()
        original = path.read_bytes()
        records, skipped = read_records(self.config)
        self.assertEqual(skipped, [])
        for record in records[:2]:
            front = render_exercise(record, self.config)['FrontField']
            self.assertIn('font-weight: 700', front)
            self.assertIn('font-style: italic', front)
            self.assertIn('color: #CC2233', front)
            self.assertIn('background-color: #FFEE99', front)
        self.assertNotIn('QuestionRuns', records[2])
        self.assertEqual(path.read_bytes(), original)

    def test_partial_word_rich_text_preserves_style_escaping_and_note_identity_in_exports(self):
        from openpyxl import Workbook
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont
        from bs4 import BeautifulSoup
        from main import html_preview
        from app.anki_exporter import note_guid
        workbook = Workbook()
        workbook.active.append([CellRichText('A ', TextBlock(InlineFont(b=True, color='FF0066CC'), 'blue'), ' <word> & friend'), 'type', '', 'hus'])
        workbook.save(self.config.section('input')['file'])
        workbook.close()
        record = read_records(self.config)[0][0]
        styled = render_exercise(record, self.config)
        plain = render_exercise({k: v for k, v in record.items() if k != 'QuestionRuns'}, self.config)
        self.assertEqual(note_guid(styled, self.config), note_guid(plain, self.config))
        self.assertEqual(styled['Identity'], plain['Identity'])
        soup = BeautifulSoup(styled['FrontField'], 'html.parser')
        blue = next(span for span in soup.select('.exercise-question span') if span.get_text() == 'blue')
        self.assertEqual(blue.get_text(), 'blue')
        self.assertIn('font-weight: 700', blue['style'])
        self.assertIn('color: #0066CC', blue['style'])
        self.assertEqual(styled['Words'], 'A blue <word> & friend')
        self.assertIn('&lt;word&gt; &amp;', styled['FrontField'])
        self.assertNotIn('<word>', styled['FrontField'])
        self.assertIn(styled['FrontField'], html_preview([styled], self.config))
        csv_path = self.directory / 'styled.csv'
        write_import_csv([styled], csv_path, self.config)
        with csv_path.open(encoding='utf-8-sig', newline='') as stream:
            imported = list(csv.reader(line for line in stream if not line.startswith('#')))[0]
        self.assertIn(styled['FrontField'], imported)
        package = self.directory / 'styled.apkg'
        generate_anki_package([styled], package, self.config, set())
        with zipfile.ZipFile(package) as archive:
            database = self.directory / 'styled.anki2'
            database.write_bytes(archive.read('collection.anki2'))
        with closing(sqlite3.connect(database)) as db:
            fields = db.execute('SELECT flds FROM notes').fetchone()[0].split('\x1f')
        import html
        self.assertEqual(html.unescape(fields[1]), styled['Words'])
        self.assertEqual(fields[3], styled['FrontField'])

    def test_formatted_gap_marker_split_across_runs_still_becomes_one_input(self):
        from openpyxl import Workbook
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont
        from bs4 import BeautifulSoup
        workbook = Workbook()
        question = CellRichText('  Jeg ', TextBlock(InlineFont(b=True), '['),
                                TextBlock(InlineFont(color='FFCC0000'), '...'), '] ',
                                TextBlock(InlineFont(i=True), 'her.\n'), 'Du […] der.  ')
        workbook.active.append([question, 'gap', '', 'bor\nbor'])
        workbook.save(self.config.section('input')['file'])
        workbook.close()
        record = read_records(self.config)[0][0]
        row = render_exercise(record, self.config)
        soup = BeautifulSoup(row['FrontField'], 'html.parser')
        self.assertEqual(len(soup.select('input.exercise-gap')), 2)
        self.assertTrue(any(span.get_text() == 'her.' for span in soup.select('.exercise-question span')))
        self.assertIn('font-style: italic', row['FrontField'])
        self.assertIn('<br>', row['FrontField'])
        self.assertNotIn('bor', row['FrontField'])
        self.assertNotIn('[...]', row['FrontField'])

    def test_formatted_prompts_follow_named_sheet_header_column_and_random_order(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = 'Exercises'
        worksheet.append(['Unused', 'Prompt', 'Type', 'Tags', 'Answer'])
        worksheet.append(['x', 'house', 'type', '', 'hus'])
        worksheet.append(['x', 'book', 'sentence', '', 'bok'])
        worksheet['B2'].font = Font(bold=True, color='FFFF0000')
        worksheet['B3'].font = Font(italic=True, color='FF0000FF')
        workbook.save(self.config.section('input')['file'])
        workbook.close()
        self.config.section('input').update(sheet='Exercises', header=True, word_column='Prompt', order='random')
        with patch('random.shuffle', side_effect=lambda rows: rows.reverse()):
            records, _ = read_records(self.config)
        self.assertEqual([r['Words'] for r in records], ['book', 'house'])
        self.assertIn('color: #0000FF', render_exercise(records[0], self.config)['FrontField'])
        self.assertIn('color: #FF0000', render_exercise(records[1], self.config)['FrontField'])

    def test_listening_format_does_not_reveal_hidden_prompt_or_change_audio_text(self):
        from openpyxl import load_workbook
        from openpyxl.styles import Font
        self.sheet([['Jeg bor her.', 'listen', '', 'I live here.']])
        workbook = load_workbook(self.config.section('input')['file'])
        workbook.active['A1'].font = Font(bold=True, color='FFFF0000')
        workbook.save(self.config.section('input')['file'])
        workbook.close()
        record = read_records(self.config)[0][0]
        self.assertEqual(record['Words'], 'Jeg bor her.')
        self.assertIn('QuestionRuns', record)
        record['Sound'] = '[sound:voice.mp3]'
        row = render_exercise(record, self.config)
        self.assertNotIn('Jeg bor her.', row['FrontField'])
        self.assertNotIn('#FF0000', row['FrontField'])
        self.assertIn('[sound:voice.mp3]', row['FrontField'])

    def test_excel_theme_indexed_colors_and_rich_decorations_are_preserved(self):
        from openpyxl import load_workbook
        from openpyxl.styles import Font, Color
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont
        self.sheet([['dark theme text', 'sentence', '', 'hei'], ['colored', 'type', '', 'hus']])
        workbook = load_workbook(self.config.section('input')['file'])
        workbook.active['A1'].font = Font(color=Color(theme=2), underline='single', strike=True)
        workbook.active['A2'] = CellRichText(TextBlock(InlineFont(color=Color(indexed=10), u='single', sz=18), 'colored'))
        workbook.save(self.config.section('input')['file'])
        workbook.close()
        records, _ = read_records(self.config)
        first = render_exercise(records[0], self.config)['FrontField']
        second = render_exercise(records[1], self.config)['FrontField']
        self.assertIn('color: #EEECE1', first)
        self.assertIn('text-decoration: underline line-through', first)
        self.assertIn('color: #FF0000', second)
        self.assertIn('font-size: 18pt', second)

    def test_unformatted_excel_keeps_shared_card_theme(self):
        self.sheet([['house', 'type', '', 'hus']])
        record = read_records(self.config)[0][0]
        self.assertNotIn('QuestionRuns', record)
        self.assertNotIn('<span style=', render_exercise(record, self.config)['FrontField'])

    def test_plain_rich_text_pieces_inherit_excel_size_instead_of_card_size(self):
        from openpyxl import Workbook
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont
        from bs4 import BeautifulSoup
        workbook = Workbook()
        workbook.active.append([CellRichText('Fill in ', TextBlock(InlineFont(b=True, sz=11), 'verbs'),
                                            ' in ', TextBlock(InlineFont(u='single', sz=11), 'present tense'),
                                            '.\n', TextBlock(InlineFont(sz=18), 'Large on purpose'), ' [...]'),
                                'gap', '', 'answer'])
        workbook.save(self.config.section('input')['file'])
        workbook.close()
        record = read_records(self.config)[0][0]
        runs = record['QuestionRuns']
        self.assertEqual(runs[0]['text'], 'Fill in ')
        self.assertEqual(runs[0]['style']['font-size'], '11pt')
        for run in runs:
            expected = '18pt' if run['text'] == 'Large on purpose' else '11pt'
            self.assertEqual(run['style']['font-size'], expected)
        row = render_exercise(record, self.config)
        question = BeautifulSoup(row['FrontField'], 'html.parser').select_one('.exercise-question')
        self.assertIn('font-size: 11pt', question.find('span', string='Fill in ')['style'])
        self.assertEqual(len(question.select('input')), 1)

    def test_fifth_meaning_column_is_optional_and_does_not_replace_answer(self):
        self.sheet([['house', 'type', 'home', 'hus', 'خانه'], ['book', 'type', '', 'bok'],
                    ['word', 'normal', '', '', '   ']])
        records, skipped = read_records(self.config)
        self.assertEqual(skipped, [])
        self.assertEqual(records[0]['Answer'], 'hus')
        self.assertEqual(records[0]['Meaning'], 'خانه')
        first = render_exercise(records[0], self.config)
        self.assertTrue(first['BackField'].startswith('<div class="gt-container custom-meaning-container">'))
        self.assertIn('<div class="gt-title">meaning</div>', first['BackField'])
        self.assertNotIn('خانه', first['FrontField'])
        self.assertIn('hus', first['BackField'])
        self.assertNotIn('custom-meaning-container', render_exercise(records[1], self.config)['BackField'])
        self.assertEqual(records[2]['Meaning'], '')

    def test_rich_personal_meaning_is_first_for_all_card_types_and_exports(self):
        from openpyxl import Workbook
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont
        from app.anki_exporter import note_values, row_back, info_html, note_guid
        workbook = Workbook()
        rich = CellRichText('My ', TextBlock(InlineFont(b=True, color='FF119944'), 'meaning'), '\n<example> & note')
        for kind in ('normal', 'type', 'sentence', 'gap', 'listen'):
            workbook.active.append(['word [...]' if kind == 'gap' else 'word', kind, '', 'answer', rich])
        workbook.save(self.config.section('input')['file'])
        workbook.close()
        records, skipped = read_records(self.config)
        self.assertEqual(skipped, [])
        rows = []
        for record in records:
            record['Info'] = info_html(self.config)
            record['Sound'] = ''
            record['FrontField'] = 'word'
            if record['Kind'] != 'normal':
                render_exercise(record, self.config)
            back = row_back(record, self.config)
            self.assertTrue(back.startswith('<div class="gt-container custom-meaning-container">'))
            self.assertIn('font-weight: 700', back)
            self.assertIn('color: #119944', back)
            self.assertIn('<br>', back)
            self.assertIn('&lt;example&gt; &amp; note', back)
            self.assertNotIn('<example>', back)
            self.assertTrue(back.endswith(record['Info']))
            self.assertNotIn('My meaning', record['FrontField'])
            changed = copy.deepcopy(record)
            changed['Meaning'] = 'updated meaning'
            self.assertEqual(note_guid(record, self.config), note_guid(changed, self.config))
            rows.append(record)
        package = self.directory / 'meaning.apkg'
        generate_anki_package(rows, package, self.config, set())
        with zipfile.ZipFile(package) as archive:
            database = self.directory / 'meaning.anki2'
            database.write_bytes(archive.read('collection.anki2'))
        with closing(sqlite3.connect(database)) as db:
            notes = db.execute('SELECT flds FROM notes').fetchall()
            models = json.loads(db.execute('SELECT models FROM col').fetchone()[0])
        self.assertEqual(len(notes), 5)
        self.assertTrue(all('custom-meaning-container' in fields for (fields,) in notes))
        self.assertEqual([f['name'] for f in models[str(self.config.section('anki')['model_id'])]['flds']], FIELDS)
        csv_path = self.directory / 'meaning.csv'
        write_import_csv(rows, csv_path, self.config)
        with csv_path.open(encoding='utf-8-sig', newline='') as stream:
            csv_rows = list(csv.reader(line for line in stream if not line.startswith('#')))
        for record, imported in zip(rows, csv_rows):
            self.assertEqual(imported[3:3 + len(note_values(record, self.config))], note_values(record, self.config))

    def test_normal_personal_meaning_precedes_dictionaries_after_reordering_or_disabling(self):
        from app.anki_exporter import note_values, note_guid, row_back, answer_template, info_html
        row = dict(Words='bok', Kind='normal', FrontField='bok', Meaning='my book',
                   lexin='<p>Lexin entry</p>', oxford='<p>Oxford entry</p>', Info=info_html(self.config))
        guid = note_guid(row, self.config)
        for sources, first in (({'lexin': {'enabled': True}, 'oxford': {'enabled': True}}, 'lexin'),
                               ({'oxford': {'enabled': True}, 'lexin': {'enabled': True}}, 'oxford'),
                               ({'lexin': {'enabled': False}}, 'lexin')):
            self.config.data['language']['source'] = 'en' if first == 'oxford' else 'no'
            # GUID is compared under the same language; ordering/toggles alone never change it.
            expected = note_guid(row, self.config)
            self.config.data['sources'] = sources
            values = note_values(row, self.config)
            self.assertTrue(values[FIELDS.index(first)].startswith('<div class="gt-container custom-meaning-container">'))
            if not any(setting['enabled'] for setting in sources.values()):
                self.assertNotIn('Lexin entry', values[FIELDS.index(first)])
            self.assertTrue(row_back(row, self.config).startswith('<div class="gt-container custom-meaning-container">'))
            self.assertIn('{{' + first + '}}', answer_template(self.config))
            self.assertTrue(answer_template(self.config).endswith('{{#Info}}{{Info}}{{/Info}}'))
            self.assertEqual(note_guid(row, self.config), expected)
        self.config.data['language']['source'] = 'no'
        self.assertEqual(note_guid(row, self.config), guid)

    def test_meaning_format_uses_fifth_column_after_named_input_column(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(['Unused', 'Input', 'Type', 'Tags', 'Answer', 'Meaning'])
        worksheet.append(['ignored', 'house', 'type', '', 'hus', 'my meaning'])
        worksheet['F2'].font = Font(bold=True, color='FF8844CC')
        workbook.save(self.config.section('input')['file'])
        workbook.close()
        self.config.section('input').update(header=True, word_column='Input')
        record = read_records(self.config)[0][0]
        self.assertEqual(record['Meaning'], 'my meaning')
        back = render_exercise(record, self.config)['BackField']
        self.assertIn('font-weight: 700', back)
        self.assertIn('color: #8844CC', back)

    def test_meaning_toggle_hides_section_for_every_kind_preserving_values_and_guids(self):
        from app.anki_exporter import note_values, row_back, note_guid
        for kind in ('normal', 'type', 'sentence', 'gap', 'listen'):
            with self.subTest(kind=kind):
                self.config.data['meaning'] = {'enabled': True}
                row = dict(Words='[...]' if kind == 'gap' else 'word', Kind=kind,
                           Answer='answer', Meaning='personal meaning', FrontField='word', Sound='', Tags=[])
                if kind != 'normal':
                    render_exercise(row, self.config)
                guid = note_guid(row, self.config)
                self.assertIn('custom-meaning-container', row_back(row, self.config))
                self.config.data['meaning']['enabled'] = False
                if kind != 'normal':
                    render_exercise(row, self.config)
                self.assertNotIn('custom-meaning-container', row_back(row, self.config))
                self.assertNotIn('personal meaning', ''.join(note_values(row, self.config)))
                self.assertEqual(row['Meaning'], 'personal meaning')
                self.assertEqual(note_guid(row, self.config), guid)

    def test_excel_fonts_include_portable_fallback_and_apkg_carries_real_bold_face(self):
        from openpyxl import Workbook
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont
        workbook = Workbook()
        workbook.active.append([CellRichText('Fill in ', TextBlock(InlineFont(rFont='Calibri', b=True, sz=11), 'verbs')),
                                'type', '', 'answer'])
        workbook.save(self.config.section('input')['file'])
        workbook.close()
        row = render_exercise(read_records(self.config)[0][0], self.config)
        self.assertIn('&quot;Calibri&quot;, &quot;Vazirmatn&quot;, Tahoma, sans-serif', row['FrontField'])
        package = self.directory / 'portable.apkg'
        generate_anki_package([row], package, self.config, set((ROOT / 'fonts').glob('*.ttf')))
        with zipfile.ZipFile(package) as archive:
            media = json.loads(archive.read('media'))
            self.assertIn('_Vazirmatn.ttf', media.values())
            self.assertIn('_Vazirmatn-Bold.ttf', media.values())
            bold_key = next(key for key, name in media.items() if name == '_Vazirmatn-Bold.ttf')
            self.assertEqual(archive.read(bold_key), (ROOT / 'fonts/_Vazirmatn-Bold.ttf').read_bytes())
            database = self.directory / 'portable.anki2'
            database.write_bytes(archive.read('collection.anki2'))
        with closing(sqlite3.connect(database)) as db:
            models = json.loads(db.execute('SELECT models FROM col').fetchone()[0])
        self.assertTrue(all('_Vazirmatn-Bold.ttf' in model['css'] for model in models.values()))

    def test_meaning_config_accepts_boolean_and_defaults_on_for_older_configs(self):
        from app.app_config import load_config
        self.sheet([['word']])
        for value in ({'enabled': False}, {'enabled': True}, None):
            data = copy.deepcopy(self.config.data)
            if value is None:
                data.pop('meaning', None)
            else:
                data['meaning'] = value
            with patch('app.app_config.tomllib.load', return_value=data):
                loaded = load_config(self.config.path)
            self.assertEqual(loaded.data['meaning']['enabled'], True if value is None else value['enabled'])
        for value in ({'enabled': 'false'}, {'enabled': 0}, {'unknown': True}, {'html': 'yes'}, 'yes'):
            data = copy.deepcopy(self.config.data)
            data['meaning'] = value
            with patch('app.app_config.tomllib.load', return_value=data), self.assertRaises(ConfigError):
                load_config(self.config.path)

    def test_html_meaning_keeps_tables_links_and_styles_but_removes_unsafe_markup(self):
        from app.exercises import meaning_html
        cell = ('<table><tr><th>Form</th><td style="color: red">reiser</td></tr></table>'
                '<a href="https://lexin.oslomet.no" onclick="bad()">Lexin</a><br><u>under</u>'
                '<script>bad()</script><span style="background: url(https://x.test/a.png)">bg</span>'
                '<a href="javascript:bad()">js</a><iframe src="https://x.test"></iframe>')
        row = {'Meaning': cell}
        self.config.data['meaning'] = {'enabled': True, 'html': True}
        output = meaning_html(row, self.config)
        for kept in ('<table>', '<th>Form</th>', 'style="color: red"', 'href="https://lexin.oslomet.no"', '<br/>', '<u>under</u>', '>js<'):
            self.assertIn(kept, output)
        for removed in ('onclick', '<script', 'url(', 'javascript:', '<iframe'):
            self.assertNotIn(removed, output)
        self.config.data['meaning']['html'] = False
        self.assertIn('&lt;table&gt;', meaning_html(row, self.config))

    def test_distributed_template_has_five_headers_and_supports_all_card_types(self):
        import warnings
        from openpyxl import load_workbook
        from zipfile import ZipFile
        template = ROOT / 'Template.xlsx'
        original = template.read_bytes()
        with warnings.catch_warnings():
            warnings.simplefilter('ignore', UserWarning)
            workbook = load_workbook(template, rich_text=True)
        worksheet = workbook.worksheets[0]
        self.assertEqual([worksheet.cell(1, i).value for i in range(1, 6)], ['Word', 'Type', 'Tag', 'Answer', 'Meaning'])
        self.assertFalse(any(worksheet.cell(row, column).value is not None
                             for row in range(2, worksheet.max_row + 1) for column in range(1, 6)))
        for i, kind in enumerate(('normal', 'type', 'sentence', 'gap', 'listen'), 2):
            for column, value in enumerate(('Jeg [...] her.' if kind == 'gap' else 'word', kind, 'sample',
                                            'bor' if kind == 'gap' else 'answer', 'personal meaning'), 1):
                worksheet.cell(i, column).value = value
        sample = self.directory / 'template-filled.xlsx'
        workbook.save(sample)
        workbook.close()
        self.config.section('input').update(file=str(sample), header=True, sheet=0, word_column=0)
        records, skipped = read_records(self.config)
        self.assertFalse(skipped)
        self.assertEqual([r['Kind'] for r in records], ['normal', 'type', 'sentence', 'gap', 'listen'])
        self.assertTrue(all(r['Meaning'] == 'personal meaning' for r in records))
        with ZipFile(template) as archive:
            self.assertIn(b'Sheet2!$A$1:$A$8', archive.read('xl/worksheets/sheet1.xml'))
        self.assertEqual(template.read_bytes(), original)


if __name__ == '__main__':
    unittest.main()
