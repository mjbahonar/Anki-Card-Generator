"""Safe Excel text formatting for exercise prompts, separate from plain identity."""
import colorsys
import html
from pathlib import Path
import xml.etree.ElementTree as ET


def _theme_colors(workbook):
    if not workbook.loaded_theme:
        return []
    root = ET.fromstring(workbook.loaded_theme)
    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    scheme = root.find('.//a:clrScheme', ns)
    colors = []
    for item in scheme if scheme is not None else []:
        node = next(iter(item), None)
        colors.append(node.get('lastClr') or node.get('val') if node is not None else None)
    # Excel theme indices put light before dark; clrScheme XML uses dark first.
    return [colors[i] for i in (1, 0, 3, 2, 4, 5, 6, 7, 8, 9, 10, 11)] if len(colors) == 12 else colors


def _color(color, theme, explicit=False):
    from openpyxl.styles.colors import COLOR_INDEX
    if color is None or color.type == 'auto':
        return None
    if color.type == 'rgb':
        value = color.rgb[-6:]
    elif color.type == 'indexed':
        value = COLOR_INDEX[color.indexed][-6:] if 0 <= color.indexed < len(COLOR_INDEX) else None
    elif color.type == 'theme':
        # Automatic/default text colors should continue to follow the card theme.
        if not explicit and color.theme == 1 and not color.tint:
            return None
        value = theme[color.theme] if 0 <= color.theme < len(theme) else None
    else:
        return None
    if not value or len(value) != 6 or any(c not in '0123456789abcdefABCDEF' for c in value):
        return None
    if color.tint:
        rgb = [int(value[i:i + 2], 16) / 255 for i in (0, 2, 4)]
        hue, light, saturation = colorsys.rgb_to_hls(*rgb)
        light = light * (1 + color.tint) if color.tint < 0 else light * (1 - color.tint) + color.tint
        value = ''.join(f'{round(v * 255):02X}' for v in colorsys.hls_to_rgb(hue, light, saturation))
    return '#' + value.upper()


def _font_style(font, theme, default=None):
    style = {}
    if font.b:
        style['font-weight'] = '700'
    elif default is None and font.b is False:
        style['font-weight'] = '400'
    if font.i:
        style['font-style'] = 'italic'
    elif default is None and font.i is False:
        style['font-style'] = 'normal'
    decorations = []
    if font.u and font.u != 'none':
        decorations.append('underline')
    if font.strike:
        decorations.append('line-through')
    if decorations:
        style['text-decoration'] = ' '.join(decorations)
    color = _color(font.color, theme, explicit=default is None)
    if color:
        style['color'] = color
    if font.sz and (default is None or font.sz != default.sz):
        style['font-size'] = f'{font.sz:g}pt'
    name = getattr(font, 'rFont', None)
    if not isinstance(name, str):
        name = getattr(font, 'name', None)
    if not isinstance(name, str):
        name = None
    if name and (default is None or name != default.name):
        # A quoted CSS string, escaped again as an HTML attribute when rendered.
        name = name.replace('\\', '\\\\').replace('"', '\\"').replace('\n', ' ').replace('\r', ' ')
        style['font-family'] = '"' + name + '", "Vazirmatn", Tahoma, sans-serif'
    if font.vertAlign in ('superscript', 'subscript'):
        style['vertical-align'] = 'super' if font.vertAlign == 'superscript' else 'sub'
    return style


def _trim_runs(runs, question):
    original = ''.join(run['text'] for run in runs)
    if original.strip() != question:
        return None
    start = len(original) - len(original.lstrip())
    end = len(original.rstrip())
    position, trimmed = 0, []
    for run in runs:
        stop = position + len(run['text'])
        text = run['text'][max(0, start - position):max(0, min(stop, end) - position)]
        if text:
            trimmed.append({'text': text, 'style': run['style']})
        position = stop
    return trimmed if any(run['style'] for run in trimmed) else None


def read_question_formats(config, records, column, field='Words', include_normal=False):
    """Read one text column. Never save or modify the user's workbook."""
    path = config.resolve(config.section('input')['file'])
    if Path(path).suffix.lower() == '.xls':
        return _xls_formats(path, config.section('input')['sheet'], records, column, field, include_normal)
    from openpyxl import load_workbook
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    workbook = load_workbook(path, data_only=True, rich_text=True)
    try:
        sheet = config.section('input')['sheet']
        worksheet = workbook.worksheets[sheet] if type(sheet) is int else workbook[sheet]
        theme, default = _theme_colors(workbook), workbook._fonts[0]
        formats = {}
        for record in records:
            if (record['Kind'] == 'normal' and not include_normal) or not record.get(field):
                continue
            cell = worksheet.cell(record['InputRow'], column + 1)
            base = _font_style(cell.font, theme, default)
            if cell.fill.patternType == 'solid':
                background = _color(cell.fill.fgColor, theme, explicit=True)
                if background:
                    base['background-color'] = background
            if isinstance(cell.value, CellRichText):
                # Plain rich-text pieces inherit the Excel cell size, not the
                # card's larger question size. Runs can still override it.
                size = cell.font.sz or default.sz
                if size:
                    base['font-size'] = f'{size:g}pt'
                runs = [{'text': part.text if isinstance(part, TextBlock) else str(part),
                         'style': {**base, **_font_style(part.font, theme)} if isinstance(part, TextBlock) else dict(base)}
                        for part in cell.value]
            else:
                runs = [{'text': str(cell.value), 'style': base}]
            trimmed = _trim_runs(runs, record[field])
            if trimmed:
                formats[record['InputRow']] = trimmed
        return formats
    finally:
        workbook.close()


def _xls_formats(path, sheet, records, column, field, include_normal):
    import xlrd
    workbook = xlrd.open_workbook(path, formatting_info=True)
    try:
        worksheet = workbook.sheet_by_index(sheet) if type(sheet) is int else workbook.sheet_by_name(sheet)
        formats = {}
        def style(font):
            result = {}
            if font.bold:
                result['font-weight'] = '700'
            if font.italic:
                result['font-style'] = 'italic'
            decorations = [name for enabled, name in ((font.underlined, 'underline'), (font.struck_out, 'line-through')) if enabled]
            if decorations:
                result['text-decoration'] = ' '.join(decorations)
            rgb = workbook.colour_map.get(font.colour_index)
            if rgb:
                result['color'] = '#' + ''.join(f'{value:02X}' for value in rgb)
            return result
        for record in records:
            if (record['Kind'] == 'normal' and not include_normal) or not record.get(field):
                continue
            row = record['InputRow'] - 1
            value = worksheet.cell_value(row, column)
            if not isinstance(value, str):
                continue
            font_index = workbook.xf_list[worksheet.cell_xf_index(row, column)].font_index
            markers = worksheet.rich_text_runlist_map.get((row, column), [])
            if not markers or markers[0][0] != 0:
                markers = [(0, font_index), *markers]
            runs = [{'text': value[start:markers[i + 1][0] if i + 1 < len(markers) else len(value)],
                     'style': style(workbook.font_list[index])} for i, (start, index) in enumerate(markers)]
            trimmed = _trim_runs(runs, record[field])
            if trimmed:
                formats[record['InputRow']] = trimmed
        return formats
    finally:
        workbook.release_resources()


def formatted_question(row, gap_replacement=None):
    """Replace gaps using plain-text offsets, even when a marker spans rich runs."""
    from app.exercises import GAP_MARKER
    text = row['Words']
    runs = row.get('QuestionRuns') or [{'text': text, 'style': {}}]
    if ''.join(run['text'] for run in runs) != text:
        runs = [{'text': text, 'style': {}}]
    def fragment(start, end):
        position, pieces = 0, []
        for run in runs:
            stop = position + len(run['text'])
            piece = run['text'][max(0, start - position):max(0, min(stop, end) - position)]
            if piece:
                escaped = html.escape(piece).replace('\n', '<br>')
                css = '; '.join(f'{key}: {value}' for key, value in run['style'].items())
                pieces.append('<span style="' + html.escape(css, quote=True) + '">' + escaped + '</span>' if css else escaped)
            position = stop
        return ''.join(pieces)
    if gap_replacement is None:
        return fragment(0, len(text))
    pieces, position = [], 0
    for match in GAP_MARKER.finditer(text):
        pieces.extend((fragment(position, match.start()), gap_replacement(match)))
        position = match.end()
    pieces.append(fragment(position, len(text)))
    return ''.join(pieces)
